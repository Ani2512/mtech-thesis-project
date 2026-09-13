from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

OUTDIR = Path(__file__).resolve().parent
FIG = OUTDIR / "figures"

INK   = RGBColor(0x0B, 0x0B, 0x0B)
BODY  = RGBColor(0x2F, 0x2E, 0x2B)
MUTED = RGBColor(0x8A, 0x88, 0x80)
SEC   = RGBColor(0x52, 0x51, 0x4E)
BLUE  = RGBColor(0x2A, 0x78, 0xD6)
ORANGE= RGBColor(0xEB, 0x68, 0x34)
RED   = RGBColor(0xB1, 0x4A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT  = RGBColor(0xF4, 0xF7, 0xFC)
LINE  = RGBColor(0xE3, 0xE1, 0xDC)
FONT  = "Helvetica Neue"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.78)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

def tf_set(tf, items, size=16, color=BODY, space=10, bullet_color=BLUE, line=1.3):
    """items: list of (text, level) or str"""
    tf.word_wrap = True
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(space)
        p.line_spacing = line
        # inline bold marked with **
        parts = text.split("**")
        for i, chunk in enumerate(parts):
            if not chunk:
                continue
            r = p.add_run(); r.text = chunk
            r.font.size = Pt(size if lvl == 0 else size - 2)
            r.font.name = FONT
            r.font.color.rgb = color if i % 2 == 0 else INK
            r.font.bold = (i % 2 == 1)
    return tf

def textbox(slide, x, y, w, h, items, **kw):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf_set(box.text_frame, items, **kw)
    return box

def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    y = Inches(0.52)
    if kicker:
        kb = s.shapes.add_textbox(M, y, W - 2*M, Inches(0.3))
        p = kb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = kicker.upper()
        r.font.size = Pt(11.5); r.font.name = FONT; r.font.bold = True
        r.font.color.rgb = MUTED
        y = Inches(0.86)
    tb = s.shapes.add_textbox(M, y, W - 2*M, Inches(0.62))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(29); r.font.name = FONT; r.font.bold = True
    r.font.color.rgb = INK
    rule = s.shapes.add_shape(1, M, y + Inches(0.68), Inches(1.05), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb = BLUE
    rule.line.fill.background(); rule.shadow.inherit = False
    PAGE[0] += 1
    footer(s, PAGE[0])
    return s

def footer(s, n):
    b = s.shapes.add_textbox(W - M - Inches(1.2), H - Inches(0.52), Inches(1.2), Inches(0.3))
    p = b.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(n)
    r.font.size = Pt(10); r.font.name = FONT; r.font.color.rgb = MUTED

def picture(s, name, top, height=None, width=None):
    path = str(FIG / name)
    if width is None:
        pic = s.shapes.add_picture(path, 0, top, height=height)
        pic.left = int((W - pic.width) / 2)
    else:
        pic = s.shapes.add_picture(path, 0, top, width=width)
        pic.left = int((W - pic.width) / 2)
    return pic

def table(s, x, y, w, rows, col_w=None, fs=12.5, head_fs=12, row_h=Inches(0.32),
          highlight=None):
    nr, nc = len(rows), len(rows[0])
    shp = s.shapes.add_table(nr, nc, x, y, w, row_h * nr)
    t = shp.table
    t.first_row = True; t.horz_banding = False
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            t.columns[i].width = Emu(int(w * cw / total))
    head_h = min(row_h, Inches(0.38))
    for ri, row in enumerate(rows):
        t.rows[ri].height = head_h if ri == 0 else row_h
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.margin_left = Inches(0.08); c.margin_right = Inches(0.08)
            c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = TINT if ri == 0 else WHITE
            base = SEC if ri == 0 else BODY
            emph = None
            if highlight and (ri, ci) in highlight:
                base = highlight[(ri, ci)]
            tf = c.text_frame
            tf.clear()
            lines = str(val).split("\n")
            for li, lineval in enumerate(lines):
                para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                para.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                para.space_after = Pt(1)
                para.line_spacing = 1.12
                for k, chunk in enumerate(lineval.split("**")):
                    if not chunk:
                        continue
                    r = para.add_run(); r.text = chunk
                    r.font.size = Pt(head_fs if ri == 0 else fs)
                    r.font.name = FONT
                    r.font.color.rgb = INK if k % 2 else base
                    r.font.bold = (ri == 0) or (k % 2 == 1) or bool(
                        highlight and (ri, ci) in highlight)
    return t

def callout(s, x, y, w, h, text, color=BLUE, fill=TINT, size=15):
    box = s.shapes.add_shape(1, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.color.rgb = color; box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.14); tf.margin_bottom = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_set(tf, [text], size=size, color=SEC, space=2)
    return box

PAGE = [1]
n = 0
# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
bar = s.shapes.add_shape(1, 0, 0, Inches(0.22), H)
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
bar.line.fill.background(); bar.shadow.inherit = False
tb = s.shapes.add_textbox(Inches(0.22), Inches(2.25), W - Inches(0.22), Inches(1.4))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Compositional Temporal\nAudio Grounding"
r.font.size = Pt(46); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = INK
p.line_spacing = 1.08
box = textbox(s, Inches(0.22), Inches(3.95), W - Inches(0.22), Inches(0.5),
        ["Finding **when** a sound happens — when the question depends on another sound"],
        size=18, color=SEC)
for para in box.text_frame.paragraphs: para.alignment = PP_ALIGN.CENTER
box = textbox(s, Inches(0.22), Inches(5.0), W - Inches(0.22), Inches(1.2),
        ["Anirudh Rangavajhala   ·   M.Tech, IIIT Dharwad",
         "12 September 2026"],
        size=14, color=MUTED, space=5)
for para in box.text_frame.paragraphs: para.alignment = PP_ALIGN.CENTER
n += 1

# ---------------------------------------------------------------- 2 problem
s = slide("A timestamp question that needs two sounds", "the task")
textbox(s, M, Inches(2.05), Inches(6.1), Inches(4),
        ["Audio LLMs are now asked **when** something happens, not just what.",
         "Today that means one sound at a time: “find the dog barking.”",
         "Real questions are **compositional** — the answer depends on a second event, on order, or on overlap."],
        size=16, space=13)
callout(s, Inches(7.3), Inches(1.95), Inches(5.25), Inches(3.6),
        "“the car horn **after** the siren”\n\n"
        "“the **second** bark”\n\n"
        "“the dog barking **while** music plays”\n\n"
        "“the bark **not followed by** a door slam”\n\n"
        "“the **next** bark after 4.0 s”", size=17)
textbox(s, M, Inches(5.85), Inches(11.8), Inches(0.9),
        ["The answer is a **set of intervals**, and it can legitimately be empty — which is itself a thing to get right."],
        size=15, color=SEC)
n += 1

# ------------------------------------------------- 3 literature review (a)
s = slide("Temporal grounding in audio LLMs", "related work  ·  1 of 2")
rows = [
 ["Work", "Contribution", "Query form"],
 ["TimeAudio (2511.11039)", "Temporal markers, time-aware encoding, segment-level token merging", "single event"],
 ["SpotSound (2604.13023)", "Suppresses hallucinated timestamps for absent events; SpotSound-Bench", "single event"],
 ["Audio-Side Time Prompt (2604.13715)", "TimePro-RL: timestamps embedded in the audio features, SFT then RL", "single event"],
 ["TEMPO (2608.29999)", "Multi-task post-training, SFT + RL across speech, music and audio", "single event"],
 ["Auto-AEG (2607.04383)", "Synthesised clips + multi-model pseudo-labels, RL; introduces AEGBench", "list all intervals of X"],
 ["AudioMap (2608.09559)", "Cloze-and-choice RL for time-aware dense audio captioning", "dense captions"],
 ["Listening with Time (2604.22245)", "Global-to-local iterative reasoning over long-form audio", "single event"],
 ["Encode Once, Decode Never (2602.10230)", "Reuses audio-LM internals to localise; 50x faster than decoding", "single event"],
 ["LA-RAG (2602.14612)", "Event-grounded QA over long audio via structured retrieval", "retrieval"],
 ["SpectCount (2606.06907)", "Counting via synthetic spectrotemporal signals", "counts"],
 ["TAG-Bench (2609.01542)", "1,750 query–recording pairs, 149.5 h, eight subsets, 7 s to 20 min", "benchmark"],
 ["DCASE 2026 Task 6", "Audio Moment Retrieval from Long Audio", "one moment"],
]
table(s, M, Inches(1.95), W - 2*M, rows, col_w=[3.5, 5.6, 2.7], fs=10.5, head_fs=10.5,
      row_h=Inches(0.335))
callout(s, M, Inches(6.35), W - 2*M, Inches(0.62),
        "Every one of these localises **one sound named on its own**. None takes a query whose answer "
        "depends on a second event.", size=13)
n += 1

# ------------------------------------------------- 4 literature review (b)
s = slide("Compositional and relational queries", "related work  ·  2 of 2")
rows = [
 ["Work", "Contribution", "What it does not do"],
 ["DAQA — Fayek & Johnson, TASLP 2020 (1911.09655)", "Programmatic before / after / ordinal / count questions; MALiMo model", "Answers are event names, yes/no or counts — never intervals"],
 ["Sridhar et al., NAACL 2025 Industry, Qualcomm (2409.06223)", "GPT-4 temporal QA from AudioSet-SL; curriculum fine-tune of LTU", "Free-text QA; no interval output, no IoU"],
 ["**TREA — Interspeech 2025 (2505.13115)**", "**Composed ESC-50, as here**; duration / ordering / counting, 600 items", "Concatenation only, so no concurrency; MCQ, not intervals"],
 ["Oncescu et al., ACM MM 2024 (2409.00851)", "Temporal ordering in text-to-audio retrieval; TempTest sets", "Retrieval, not localisation"],
 ["PolyBench — Interspeech 2026 (2603.05128)", "Counting, classification, detection, concurrency, duration on polyphonic audio", "Multiple choice; no conditional interval sets"],
 ["CompA-order — ICLR 2024 (2310.08753)", "400 audio–caption matching instances probing event order", "Matching, not grounding"],
 ["STAR-Bench (2510.24693), MMAU (2410.19168)", "Segment reordering, spatial localisation; 27 tasks, 10k clips", "Multiple choice; no grounding"],
 ["CoMET-Bench (2606.15320) — **video**", "Conditional multi-event grounding, 4 temporal + 3 spatial conditions,\nRejection-F1, CoMET-Agent", "Video only; no superposition-based “while”"],
 ["CompSTVG (2608.30584) — **video**", "Scene-graph synthetic curriculum plus curriculum RL", "Video only"],
 ["MUSEG (2505.20715),\nOne-to-Many TG (2606.06294) — **video**", "Multi-segment grounding; RL with completeness rewards", "Video only"],
]
table(s, M, Inches(1.92), W - 2*M, rows, col_w=[4.1, 4.4, 3.3], fs=9.5, head_fs=10,
      row_h=Inches(0.42))
callout(s, M, Inches(6.44), W - 2*M, Inches(0.5),
        "**TREA is the closest neighbour**: same ESC-50 composition. It concatenates, so concurrency "
        "cannot arise, and its answers are multiple choice.", size=12)
n += 1

# ---------------------------------------------------------------- 5 the gap
s = slide("The gap the two halves leave open", "positioning")
rows = [["", "Interval-set output", "Condition on a second event", "Audio"],
        ["Audio grounding work (TAG-Bench, TEMPO, Auto-AEG, …)", "yes", "no", "yes"],
        ["Audio relational QA (DAQA, Sridhar et al., PolyBench)", "no", "yes", "yes"],
        ["Video compositional grounding (CoMET-Bench, CompSTVG)", "yes", "yes", "no"],
        ["**This project**", "**yes**", "**yes**", "**yes**"]]
hl = {(4,1): BLUE, (4,2): BLUE, (4,3): BLUE}
table(s, M, Inches(2.15), W - 2*M, rows, col_w=[5.6, 2.1, 2.9, 1.2], fs=13,
      row_h=Inches(0.52), highlight=hl)
textbox(s, M, Inches(4.78), Inches(11.8), Inches(1.5),
        ["Each half of the task has a neighbour; the intersection is empty in audio. The video precedent "
         "(CoMET-Bench, 2026) shows the task shape is real and worth doing.",
         "**Cited, not claimed:** DAQA and Sridhar et al. are the query-generation precedents — the question "
         "templates are not new. The framing is a port of CoMET-Bench from video, and reviewers will say so."],
        size=13.5, space=10)
callout(s, M, Inches(6.42), W - 2*M, Inches(0.66),
        "The defence: mixing-made **while** conditions (impossible in video), rejection queries, "
        "the ceiling diagnostic, and the miss/false-alarm asymmetry.", size=13.5)
n += 1

# ---------------------------------------------------------------- 4 what was built
s = slide("What was built", "method")
textbox(s, M, Inches(2.05), Inches(5.9), Inches(4.4),
        ["**A benchmark generator.** Sounds are composed onto a known timeline, so ground truth is exact — no annotation, no label noise.",
         "**Eight query types**, including rejection queries whose correct answer is “no such event”.",
         "**Metrics** beyond f1: union-IoU, Hungarian matching, count accuracy, under-report rate, centre error, duration ratio.",
         "**Five arms** from one codebase, 40 passing tests."],
        size=15.5, space=12)
rows = [["Type", "Query", "n (test)"],
        ["PLAIN", "the dog bark", "96"],
        ["ORDINAL", "the second bark", "96"],
        ["BEFORE / AFTER", "the bark after the siren", "89 / 89"],
        ["NEXT_AFTER", "the next bark after 4.0 s", "89"],
        ["WHILE", "the bark while music plays", "96"],
        ["NOT_FOLLOWED", "a bark not followed by a slam", "96"],
        ["ABSENT", "(rejection queries)", "48"]]
table(s, Inches(7.0), Inches(2.05), Inches(5.55), rows, col_w=[2.4, 4.4, 1.3], fs=11.5, row_h=Inches(0.37))
callout(s, Inches(7.0), Inches(5.55), Inches(5.55), Inches(0.72),
        "**4404 queries** over 300 composed ESC-50 clips.", size=14)
n += 1

# ---------------------------------------------------------------- 5 backbone
# ---------------------------------------------------------------- architecture
s = slide("Architecture, as it runs today", "system")
picture(s, "fig5_architecture.png", Inches(1.8), height=Inches(4.55))
callout(s, M, Inches(6.5), W - 2*M, Inches(0.55),
        "Three layers. **Build** makes the audio and therefore knows the answer. **Answer** is five arms from one "
        "codebase. **Score** is one metrics module, so every arm is measured the same way.", size=12.5)
n += 1

s = slide("Only one of the two candidate backbones can ground at all", "phase 1")
picture(s, "fig1_backbone.png", Inches(2.0), height=Inches(3.7))
textbox(s, M, Inches(5.95), Inches(11.8), Inches(1.1),
        ["Qwen2-Audio scores **0.094** on plain grounding and **exactly 0.000** on every conditional type — it emits well-formed intervals that land nowhere near the event. SpotSound (2604.13023, §4.1) reports the same failure mode: “syntactically complete time windows … resulting in zero overlap with the ground truth.”",
         "**Qwen2.5-Omni is the backbone.** It finds the sound, then drops the qualifier attached to it."],
        size=14, space=7)
n += 1

# ---------------------------------------------------------------- 6 ordering
s = slide("The hard part is not “conditions” — it is relating two events", "phase 1")
rows = [["", "Handled", "Collapses"],
        ["Qwen2.5-Omni, f1@0.5", "BEFORE 0.379    ORDINAL 0.333",
         "WHILE 0.152    AFTER 0.146\nNOT_FOLLOWED 0.095    NEXT_AFTER 0.091"]]
table(s, M, Inches(2.15), W - 2*M, rows, col_w=[2.8, 4.4, 4.6], fs=13, row_h=Inches(0.85))
textbox(s, M, Inches(3.55), Inches(11.8), Inches(2.2),
        ["ORDINAL needs only counting inside one event type. BEFORE needs one comparison against a single reference.",
         "The four that collapse all require **relating the target to a second event and then selecting among the matches**. That is the specific capability that is missing.",
         "Supporting signal: on NOT_FOLLOWED the model under-reports on **69%** of queries — the highest of any type — consistent with collapsing a set of matches down to one."],
        size=16, space=13)
n += 1

# ---------------------------------------------------------------- 7 ceiling
s = slide("Is the model bad at conditions, or bad at hearing?", "the diagnostic")
textbox(s, M, Inches(2.0), Inches(11.8), Inches(0.6),
        ["Answer the condition **for** the model: ground each sound separately, then apply the condition with the same predicates that define the ground truth. Then degrade the grounder on purpose."],
        size=15, color=SEC, space=6)
rows = [["Grounding quality given to the composer", "f1@0.5, all 4404 queries"],
        ["Perfect", "1.000"],
        ["Boundaries jittered ±0.5 s", "0.925"],
        ["25% of events missed", "0.650"],
        ["25% spurious detections", "0.937"],
        ["Degraded until PLAIN matches Qwen2.5-Omni", "0.263"],
        ["Qwen2.5-Omni answering directly", "0.243"]]
hl = {(5,1): BLUE, (6,1): BLUE}
table(s, M, Inches(2.9), Inches(8.2), rows, col_w=[5.6, 2.6], fs=13.5, row_h=Inches(0.42), highlight=hl)
callout(s, Inches(9.4), Inches(2.9), Inches(3.15), Inches(2.94),
        "Perfect events give **1.000** on every condition type.\n\n"
        "So the conditions are trivial.\n\n"
        "**Grounding quality is the binding constraint** — essentially the whole gap from 1.000 to 0.26.", size=14)
n += 1

# ---------------------------------------------------------------- 8 prediction
s = slide("Which produced a falsifiable prediction", "the diagnostic")
textbox(s, M, Inches(2.1), Inches(11.8), Inches(0.7),
        ["At the model's real grounding quality the two approaches tie overall — 0.263 against 0.243 — but the **per-type split is not a tie**, and that is testable:"],
        size=15.5, color=SEC, space=8)
rows = [["Predicted: decomposition wins", "Predicted: direct prompting wins"],
        ["AFTER, NEXT_AFTER, WHILE, NOT_FOLLOWED\n— the four needing a second event",
         "ORDINAL, BEFORE\n— the two the model already handles,\nwhere decomposition pays to ground a\nreference it does not need"]]
table(s, M, Inches(3.1), W - 2*M, rows, col_w=[5.9, 5.9], fs=14, row_h=Inches(1.25))
callout(s, M, Inches(5.85), W - 2*M, Inches(0.95),
        "This prediction was made on **CPU, from simulation only**, before any of it was run on a GPU. "
        "The next slide is the test.", size=15)
n += 1

# ---------------------------------------------------------------- 9 HEADLINE
s = slide("Tested on the real model: four of six predictions held", "result — 11 September 2026, held-out test split")
picture(s, "fig2_direct_vs_decompose.png", Inches(1.95), height=Inches(3.85))
textbox(s, M, Inches(6.05), Inches(11.8), Inches(1.0),
        ["Decomposition rescues the relational types as predicted — **NEXT_AFTER +57%, NOT_FOLLOWED +47%** — and loses BEFORE as predicted. WHILE and ORDINAL went the other way.",
         "Overall the arms nearly cancel (0.207 vs 0.218). The per-type structure, not the average, is the finding — and it is what a per-type hybrid should exploit."],
        size=13.5, space=6)
n += 1

# ---------------------------------------------------------------- 10 asymmetry
s = slide("Which grounding error actually costs us?", "the method")
picture(s, "fig3_asymmetry.png", Inches(2.0), height=Inches(3.55))
textbox(s, M, Inches(5.75), Inches(11.8), Inches(1.2),
        ["A **missed** event removes a correct interval *and* can corrupt every query that references it — lose the siren and “the horn after the siren” becomes unanswerable. A **spurious** event usually adds one wrong interval and leaves the rest intact.",
         "Cost per unit miss rate **1.411**; per unit false-alarm rate **0.252**. Asymmetry **5.60×**."],
        size=14, space=7)
n += 1

# ---------------------------------------------------------------- 11 beta
s = slide("So the objective is measured, not chosen", "the method")
big = s.shapes.add_textbox(M, Inches(2.2), Inches(5.4), Inches(1.5))
p = big.text_frame.paragraphs[0]
r = p.add_run(); r.text = "β = 2.37"
r.font.size = Pt(62); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = BLUE
textbox(s, M, Inches(3.45), Inches(5.6), Inches(2.5),
        ["β² = slope(miss) / slope(false alarm) = 1.411 / 0.252 = **5.60**",
         "β = √5.60 = **2.37**  —  read off the curves, not tuned",
         "F_β = (1 + β²) · P · R / (β² · P + R)",
         "Pinned by a unit test so it cannot quietly become a hyperparameter."],
        size=13.5, space=7)
rows = [["Prediction against 3 gold intervals", "f1", "F-2.37"],
        ["misses one", "0.800", "0.702"],
        ["invents one", "0.857", "0.952"]]
table(s, Inches(6.9), Inches(2.35), Inches(5.65), rows, col_w=[3.3, 1.2, 1.2], fs=14, row_h=Inches(0.46))
callout(s, Inches(6.9), Inches(4.2), Inches(5.65), Inches(1.5),
        "Under f1 the two are nearly equivalent. Under the **measured** beta they are not — "
        "and the training signal follows: the rejected side of a preference pair always under-detects.", size=14)
callout(s, M, Inches(5.95), W - 2*M, Inches(0.9),
        "This quantity is only measurable on a benchmark where conditions depend on a second event. "
        "**New task → new measurable quantity → objective derived from it.**", size=15)
n += 1

# ---------------------------------------------------------------- 12 negative
s = slide("A negative result: the decoder did not survive contact", "result")
picture(s, "fig4_decoding.png", Inches(1.95), height=Inches(3.6))
textbox(s, M, Inches(5.75), Inches(11.8), Inches(1.2),
        ["Sampling k times and keeping intervals that appear in at least 2 samples reached **0.951** on a simulated grounder with independent noise. On the real model it scores **0.150**, below single-sample decoding.",
         "The simulation's independence assumption was the load-bearing one. **The β = 2.37 finding is untouched** — it is a property of the benchmark, not of the decoder — but this mechanism is reported as refuted."],
        size=13.5, space=6)
n += 1

# ---------------------------------------------------------------- 13 status
s = slide("Where it stands", "status, 12 September 2026")
rows = [["Arm", "What it is", "Status"],
        ["A  direct prompting", "Ask the model the conditional query", "Run — 0.207"],
        ["B  decompose-and-combine", "Two plain calls, condition applied locally", "Run — 0.218"],
        ["D  hybrid", "Per-type choice between A and B", "Blocked — no validation split was scored"],
        ["C  QLoRA, text timestamps", "Fine-tune the thinker, encoder frozen", "Never trained — CUDA OOM"],
        ["E  QLoRA, atomic time tokens", "One token per quantised time (per TEMPO)", "Never trained — CUDA OOM"]]
hl = {(3,2): RED, (4,2): RED, (5,2): RED, (1,2): BLUE, (2,2): BLUE}
table(s, M, Inches(2.1), W - 2*M, rows, col_w=[3.3, 5.0, 3.5], fs=12.5, row_h=Inches(0.5), highlight=hl)
textbox(s, M, Inches(5.35), Inches(11.8), Inches(1.4),
        ["A 5 h 34 m run on two T4s completed the inference arms and **failed both fine-tuning arms**: backward pass out of memory on a 14.6 GB card. Neither has ever trained.",
         "Fixes identified: gradient checkpointing and shorter audio windows, or one larger GPU. Two smaller bugs — a missing import that killed the precision sweep, and an unscored validation split that silently turned the hybrid into a copy of arm A."],
        size=14, space=8)
n += 1

# ---------------------------------------------------------------- 14 limits
s = slide("What this evidence does not yet support", "limitations")
textbox(s, M, Inches(2.05), Inches(11.8), Inches(3.8),
        ["**Composed audio only.** Sounds mixed onto a synthetic timeline. Nothing here generalises to real recordings yet — that is the critical path, not an afterthought.",
         "**One prompt, one model.** No prompt-sensitivity study. TAG-Bench notes that an explicit “enumerate all intervals” instruction is an untested intervention; it should be baseline #1.",
         "**No fine-tuned arm.** Every number on these slides is zero-shot or training-free.",
         "**“Cannot ground” is benchmark-specific.** Another paper (2607.25355) measures Qwen2-Audio at 0.3653 mIoU zero-shot on its own task. Different data, different metric, 4-bit here — so the claim is about this benchmark and this prompt, not the model in general.",
         "**f1@0.5 conflates two failures** — wrong place and wrong duration. Centre error and duration ratio are reported alongside it for that reason.",
         "**The 5.60× asymmetry is measured on this query mix.** On another distribution, re-derive β rather than importing this number."],
        size=15, space=13)
n += 1

# ---------------------------------------------------------------- 15 plan
s = slide("Six-month plan", "next")
rows = [["Month", "Work", "Why now"],
        ["1", "Fix the OOM, train arms C and E; put the benchmark and the\ndecomposition diagnosis on arXiv", "TAG-Bench (Alibaba) appeared 2 Sep and is one step away"],
        ["2–3", "**Real recordings** — derive conditional queries from DESED /\nAudioSet-strong, hand-verify ~300 clips, rerun every arm", "The critical path. Starts in month 2, not month 4"],
        ["4", "β as a verifiable GRPO reward rather than offline preference pairs;\nrerun recall-biased decoding with a real model", "The offline-pair and voting forms are now known to be weak"],
        ["5–6", "Writing and defence", ""]]
table(s, M, Inches(2.05), W - 2*M, rows, col_w=[1.0, 6.4, 4.4], fs=12, row_h=Inches(0.62))
callout(s, M, Inches(5.95), W - 2*M, Inches(1.0),
        "Out of scope, stated as future work: large-scale annotation, long-form audio (where a flat timestamp "
        "vocabulary grows with duration), and whether the 5.60× asymmetry also holds in **video** — the most "
        "interesting extension, since it would generalise the claim beyond audio.", size=14)
n += 1

# ---------------------------------------------------------------- 16 contributions
s = slide("Contributions, stated honestly", "summary")
rows = [["Novel here", "Not novel — cited, not claimed"],
        ["The task and benchmark: interval-set answers to\nconditional temporal queries over audio, with exact\nprogrammatic ground truth and rejection queries",
         "Atomic timestamp tokens — TEMPO, TimeAudio"],
        ["The ceiling diagnostic that separates “conditions are\nhard” from “grounding is bad”, and the per-type\nprediction it generates",
         "Decompose-and-combine — CoMET-Agent, in video"],
        ["The **5.60× miss/false-alarm asymmetry** and the\nobjective derived from it — measurable only on a\nbenchmark where conditions depend on a second event",
         "LoRA fine-tuning for grounding — standard\n\nTask framing is a port of CoMET-Bench from video;\nthe defences are the mixing-made WHILE conditions,\nthe ceiling diagnostic, and the asymmetry"]]
table(s, M, Inches(2.1), W - 2*M, rows, col_w=[5.9, 5.9], fs=12.5, row_h=Inches(1.15))
n += 1

out = OUTDIR / "Compositional_Temporal_Audio_Grounding_2026-09-12.pptx"
prs.save(str(out))
print("saved:", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
